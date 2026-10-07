from pathlib import Path
import pymupdf


def pdf_to_pptx(input_path: Path, output_path: Path) -> Path:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for PDF to PPTX conversion") from exc

    source = pymupdf.open(input_path)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    try:
        for page in source:
            slide = presentation.slides.add_slide(blank)
            pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            image_path = output_path.with_name(f".{output_path.stem}_page_{page.number}.png")
            pix.save(image_path)
            try:
                slide.shapes.add_picture(image_path, 0, 0, width=presentation.slide_width, height=presentation.slide_height)
            finally:
                image_path.unlink(missing_ok=True)
        presentation.save(output_path)
    finally:
        source.close()

    return output_path
